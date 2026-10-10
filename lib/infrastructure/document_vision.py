"""Bounded local rasterization for explicitly enabled document vision models."""
from __future__ import annotations

import base64
import hashlib
import io
from pathlib import Path
from zipfile import ZipFile

from PIL import Image, ImageOps

from lib.domain.document_parser import supports_vision

IMAGE_EXTENSIONS = frozenset({".png", ".jpg", ".jpeg", ".webp"})
MAX_PAGES = 200
MAX_IMAGE_PIXELS = 40_000_000
MAX_EXPANDED_DOCX_BYTES = 100 * 1024 * 1024


def require_vision_model(root: Path, binding: dict) -> None:
    from lib.bootstrap.backends import backend_application
    endpoint = backend_application(root).merged_backends().get(binding["backend"]) or {}
    if not supports_vision(endpoint, binding["model"]):
        raise ValueError("document_model_vision_not_confirmed")


def _image(data: bytes, location: str) -> dict:
    try:
        with Image.open(io.BytesIO(data)) as opened:
            if opened.width * opened.height > MAX_IMAGE_PIXELS:
                raise ValueError("document_image_too_large")
            image = ImageOps.exif_transpose(opened).convert("RGB")
            image.thumbnail((2048, 2048))
            output = io.BytesIO()
            image.save(output, format="PNG")
    except (Image.DecompressionBombError, Image.DecompressionBombWarning, OSError):
        raise ValueError("invalid_document_image") from None
    payload = output.getvalue()
    return {"location": location, "image_sha256": hashlib.sha256(payload).hexdigest(),
            "image": "data:image/png;base64," + base64.b64encode(payload).decode("ascii")}


def visual_parts(path: Path, *, source_processing_version: int = 2):
    """Yield a PDF page, standalone image, or ordered DOCX text/image block."""
    path = Path(path)
    if path.suffix in IMAGE_EXTENSIONS:
        yield _image(path.read_bytes(), "image:1")
    elif path.suffix == ".pdf":
        try:
            import pypdfium2 as pdfium
        except ImportError:
            raise ValueError("document_pdf_renderer_unavailable") from None
        try:
            document = pdfium.PdfDocument(str(path))
        except pdfium.PdfiumError as error:
            if source_processing_version < 2:
                raise
            raise ValueError("document_pdf_parse_failed") from error
        try:
            if len(document) > MAX_PAGES:
                raise ValueError("document_vision_page_limit")
            for index in range(len(document)):
                page = document[index]
                try:
                    width, height = page.get_size()
                    if min(width, height) <= 0 or max(width, height) > 100_000:
                        raise ValueError("document_pdf_page_invalid")
                    bitmap = page.render(scale=min(2.0, 2048 / max(width, height)))
                    try:
                        image = bitmap.to_pil()
                        output = io.BytesIO()
                        image.save(output, format="PNG")
                        yield _image(output.getvalue(), f"page:{index + 1}")
                    finally:
                        bitmap.close()
                finally:
                    page.close()
        except pdfium.PdfiumError as error:
            if source_processing_version < 2:
                raise
            raise ValueError("document_pdf_parse_failed") from error
        finally:
            document.close()
    elif path.suffix == ".docx":
        from lib.infrastructure.docx_document import block_text, open_document
        if source_processing_version >= 2:
            document = open_document(path, max_expanded_bytes=MAX_EXPANDED_DOCX_BYTES)
        else:
            # Preserve older run behavior and all valid source block identities.
            with ZipFile(path) as archive:
                entries = archive.infolist()
                if len(entries) > 5000 or sum(item.file_size for item in entries) > MAX_EXPANDED_DOCX_BYTES:
                    raise ValueError("document_docx_expansion_limit")
            import docx
            document = docx.Document(path)
        images = 0
        for index, block in enumerate(document.element.body, 1):
            text = block_text(block, document)
            if text.strip():
                yield {"location": f"block:{index}", "text": text}
            for blip in block.xpath(".//a:blip"):
                relationship = blip.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed")
                if not relationship:
                    raise ValueError("document_external_image_unsupported")
                images += 1
                if images > MAX_PAGES:
                    raise ValueError("document_vision_page_limit")
                part = document.part.related_parts[relationship]
                yield _image(part.blob, f"block:{index}:image:{images}")
    else:
        raise ValueError("unsupported_vision_document")


def visual_part(path: Path, location: str) -> dict | None:
    """Rasterize only a requested pinned page, not all preceding PDF pages."""
    path = Path(path)
    if path.suffix.lower() == ".pdf" and location.startswith("page:"):
        try:
            index = int(location.split(":", 1)[1]) - 1
        except (ValueError, TypeError):
            raise ValueError("document_source_page_invalid") from None
        try:
            import pypdfium2 as pdfium
        except ImportError:
            raise ValueError("document_pdf_renderer_unavailable") from None
        document = pdfium.PdfDocument(str(path))
        try:
            if len(document) > MAX_PAGES:
                raise ValueError("document_vision_page_limit")
            if not 0 <= index < len(document):
                raise ValueError("document_source_page_invalid")
            page = document[index]
            try:
                width, height = page.get_size()
                if min(width, height) <= 0 or max(width, height) > 100_000:
                    raise ValueError("document_pdf_page_invalid")
                bitmap = page.render(scale=min(2.0, 2048 / max(width, height)))
                try:
                    output = io.BytesIO()
                    bitmap.to_pil().save(output, format="PNG")
                    return _image(output.getvalue(), location)
                finally:
                    bitmap.close()
            finally:
                page.close()
        finally:
            document.close()
    if path.suffix.lower() not in IMAGE_EXTENSIONS | {".docx"}:
        return None
    for part in visual_parts(path):
        if part["location"] == location:
            return part if "image" in part else None
    return None
