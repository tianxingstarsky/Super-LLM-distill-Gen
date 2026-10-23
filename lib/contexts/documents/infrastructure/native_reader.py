"""Explicit local readers. Optional PDF/DOCX dependencies load only on use."""
from __future__ import annotations
import pathlib
from ..domain.formats import SUPPORTED_TEXT_EXTENSIONS

SUPPORTED_EXTS = SUPPORTED_TEXT_EXTENSIONS

def _decode_text(data: bytes) -> str:
    """UTF-8 优先；失败再按 GB18030（中文 Windows 常见）解码。

    旧实现用 ``errors="replace"``：GBK 文档会被静默替换成 U+FFFD 乱码并写进
    "无损" CPT 语料。这里改为显式解码，两种编码都不成立时抛错而不是产出坏数据。
    """
    for encoding in ("utf-8", "gb18030"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise UnicodeError("invalid_text_encoding_utf8_gb18030")


def import_text(path: str | pathlib.Path, *, source_processing_version: int = 2) -> str:
    if type(source_processing_version) is not int or source_processing_version not in {1, 2}:
        raise ValueError("invalid_source_processing_version")
    path = pathlib.Path(path)
    ext = path.suffix.lower()
    if ext not in SUPPORTED_EXTS:
        raise ValueError(f"不支持的文件类型 {ext}（支持 {sorted(SUPPORTED_EXTS)}）")
    if ext in (".md", ".txt"):
        return _decode_text(path.read_bytes())
    if ext in (".tex", ".latex"):
        from ..domain.latex import parse_latex

        return parse_latex(_decode_text(path.read_bytes()))
    if ext == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        return "\n".join((page.extract_text() or "") for page in reader.pages)
    if ext == ".docx":
        import docx

        if source_processing_version >= 2:
            from lib.infrastructure.docx_document import document_text, open_document

            return document_text(open_document(path))
        # An older submitted run keeps the same chunks and checkpoint identities.
        document = docx.Document(str(path))
        parts = [p.text for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                parts.append("\t".join(cell.text for cell in row.cells))
        return "\n".join(parts)
    raise AssertionError("unreachable")



class LocalDocumentTextReader:
    def read_text(self, source_path: str, *, source_processing_version: int = 2) -> str:
        return import_text(source_path, source_processing_version=source_processing_version)
