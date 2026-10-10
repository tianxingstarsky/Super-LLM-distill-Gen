"""Vision documents use actual local pages and explicit capability declarations, offline."""
import importlib.util
import json
from pathlib import Path

from PIL import Image
import pytest
import yaml

from lib.domain.document_parser import supports_vision
from lib.infrastructure import training_workflow as engine
from lib.infrastructure.document_vision import visual_parts
from lib.infrastructure.workflow_stream_journal import read_streams


TEXT = "Before starting the equipment, check the power cable. Switch off the power before maintenance. Record the result after every inspection."


def config(root, *, confirmed=True, api_format="chat"):
    folder = root / "configs"
    folder.mkdir(exist_ok=True)
    (folder / "preferences.yaml").write_bytes((Path(__file__).resolve().parents[1] / "configs/preferences.yaml").read_bytes())
    endpoint = {"base_url": "https://offline.invalid/v1", "models": ["writer"],
                "api_format": api_format, "api_key_env": "VISION_TEST_KEY"}
    if confirmed:
        endpoint["model_capabilities"] = {"writer": {"vision": True}}
    (folder / "backends.yaml").write_text(yaml.safe_dump({"backends": {"vision": endpoint}}), encoding="utf-8")


def parser():
    return {"mode": "vision", "binding": {"backend": "vision", "model": "writer",
            "context_window_tokens": 131072, "max_output_tokens": 32768}}


@pytest.mark.parametrize("name, content, reason", [
    ("broken.docx", b"not a document package", "document_docx_parse_failed"),
    ("broken.pdf", b"%PDF-1.7\nnot a PDF body", "document_pdf_parse_failed"),
    ("broken.png", b"not a picture", "invalid_document_image"),
])
def test_corrupt_visual_source_does_not_abort_other_documents(tmp_path, name, content, reason):
    config(tmp_path)
    invalid = tmp_path / name
    invalid.write_bytes(content)
    valid = tmp_path / "valid.txt"
    valid.write_text(TEXT, encoding="utf-8")
    output = tmp_path / "output"
    run_id = engine.create_run(output, sources=[invalid, valid], targets=["cpt"],
                               settings_root=tmp_path, document_parser=parser())
    reader = Vision()
    run = engine.Workflow(output, run_id, tmp_path, generator=reader)
    state = run.execute()
    assert state["status"] == "needs_attention"
    assert reader.calls == []
    assert engine.verify_artifacts(run.path)["counts"]["cpt"] == 1
    inputs = engine.read_json(run.path / "input_records.json")
    assert any(row.get("reason") == reason and row["source_name"] == name for row in inputs)


def image(root, name="scan.png"):
    path = root / name
    Image.new("RGB", (120, 80), "blue").save(path)
    return path


class Vision:
    model = "writer"

    def __init__(self, *, uncertain=False):
        self.calls = []
        self.usage = {"calls": 0}
        self.uncertain = uncertain

    def chat(self, messages, **options):
        self.calls.append((messages, options))
        self.usage["calls"] += 1
        assert messages[1]["content"][1]["type"] == "image_url"
        return json.dumps({"text": TEXT, "uncertain": self.uncertain})


def test_model_name_and_role_do_not_imply_vision_support():
    assert not supports_vision({"models": ["vision-vl-multimodal"], "roles": ["vision"]}, "vision-vl-multimodal")
    assert not supports_vision({"model_capabilities": {"writer": {"vision": "true"}}}, "writer")
    assert supports_vision({"model_capabilities": {"writer": {"vision": True}}}, "writer")
    assert not supports_vision({"model_capabilities": {"different": {"vision": True}}}, "writer")


def test_unconfirmed_model_and_native_image_are_rejected_before_run_creation(tmp_path):
    source = image(tmp_path)
    config(tmp_path, confirmed=False)
    output = tmp_path / "output"
    with pytest.raises(ValueError, match="document_model_vision_not_confirmed"):
        engine.create_run(output, sources=[source], targets=["cpt"], settings_root=tmp_path, document_parser=parser())
    with pytest.raises(ValueError, match="document_image_requires_vision_parser"):
        engine.create_run(output, sources=[source], targets=["cpt"])
    assert not output.exists()


def test_image_workflow_exports_real_transcription_provenance_and_reuses_checkpoint(tmp_path):
    config(tmp_path)
    source = image(tmp_path)
    output = tmp_path / "output"
    run_id = engine.create_run(output, sources=[source], targets=["cpt"], settings_root=tmp_path,
                               document_parser=parser())
    client = Vision()
    run = engine.Workflow(output, run_id, tmp_path, generator=client)
    assert run.recipe["node_models"]["ingest"]["vision"]["model"] == "writer"
    assert run.recipe["endpoint_pins"]["ingest"]["vision"]["backend"] == "vision"
    state = run.execute()
    assert state["status"] == "completed"
    assert state["usage"]["vision"]["calls"] == 1
    records = engine.read_json(run.path / "artifacts/cpt.records.json")
    assert records[0]["text"] == TEXT
    assert records[0]["source_name"] == "scan.png"
    assert records[0]["source_location"]["record"] == "image:1"
    assert records[0]["evidence_level"] == "model_transcribed_visual_source"
    assert records[0]["document_reading"]["image_sha256"]
    engine.verify_artifacts(run.path)
    # A fresh worker can reuse the durable model result without a new API call.
    restored = engine.Workflow(output, run_id, tmp_path, generator=client)
    list(restored.iter_source_units(restored.recipe["sources"][0]))
    assert len(client.calls) == 1
    assert client.calls[0][1]["max_tokens"] == 32768


def test_uncertain_visual_reading_is_quarantined_not_published(tmp_path):
    config(tmp_path)
    output = tmp_path / "output"
    run_id = engine.create_run(output, sources=[image(tmp_path)], targets=["cpt"], settings_root=tmp_path,
                               document_parser=parser())
    run = engine.Workflow(output, run_id, tmp_path, generator=Vision(uncertain=True))
    state = run.execute()
    assert state["status"] == "needs_attention"
    assert engine.verify_artifacts(run.path)["counts"]["cpt"] == 0
    inputs = engine.read_json(run.path / "input_records.json")
    assert inputs[0]["reason"] == "document_vision_uncertain"


def test_revoked_capability_blocks_execution_of_queued_recipe(tmp_path):
    config(tmp_path)
    run_id = engine.create_run(tmp_path / "output", sources=[image(tmp_path)], targets=["cpt"],
                               settings_root=tmp_path, document_parser=parser())
    config(tmp_path, confirmed=False)
    with pytest.raises(ValueError, match="document_model_vision_not_confirmed"):
        engine.Workflow(tmp_path / "output", run_id, tmp_path, generator=Vision())


def test_pdf_is_rendered_page_by_page_with_real_rasters_and_locations(tmp_path):
    source = tmp_path / "scan.pdf"
    first, second = Image.new("RGB", (120, 90), "blue"), Image.new("RGB", (120, 90), "red")
    first.save(source, save_all=True, append_images=[second])
    pages = list(visual_parts(source))
    assert [part["location"] for part in pages] == ["page:1", "page:2"]
    assert all(part["image"].startswith("data:image/png;base64,") for part in pages)
    assert pages[0]["image_sha256"] != pages[1]["image_sha256"]


@pytest.mark.parametrize("phase", ["load", "render", "bitmap"])
@pytest.mark.parametrize("known_pdf_error", [True, False])
@pytest.mark.parametrize("source_processing_version", [1, 2])
def test_pdf_page_failures_release_resources_and_only_normalize_known_errors(
        tmp_path, monkeypatch, phase, known_pdf_error, source_processing_version):
    import pypdfium2 as pdfium

    error_type = pdfium.PdfiumError if known_pdf_error else RuntimeError
    closed = []

    class Bitmap:
        def to_pil(self):
            raise error_type("offline page bitmap failure")

        def close(self):
            closed.append("bitmap")

    class Page:
        def get_size(self):
            return 100, 100

        def render(self, **options):
            if phase == "render":
                raise error_type("offline page rendering failure")
            return Bitmap()

        def close(self):
            closed.append("page")

    class Document:
        def __len__(self):
            return 1

        def __getitem__(self, index):
            if phase == "load":
                raise error_type("offline page loading failure")
            return Page()

        def close(self):
            closed.append("document")

    monkeypatch.setattr(pdfium, "PdfDocument", lambda *_: Document())
    normalized = known_pdf_error and source_processing_version == 2
    with pytest.raises(ValueError if normalized else error_type) as failure:
        list(visual_parts(tmp_path / "mock.pdf", source_processing_version=source_processing_version))
    if normalized:
        assert str(failure.value) == "document_pdf_parse_failed"
        assert type(failure.value.__cause__) is pdfium.PdfiumError
    else:
        assert "offline page" in str(failure.value)
    assert closed == {"load": ["document"], "render": ["page", "document"],
                      "bitmap": ["bitmap", "page", "document"]}[phase]


def test_docx_preserves_text_and_embedded_image_in_body_order(tmp_path):
    import docx
    source = tmp_path / "guide.docx"
    document = docx.Document()
    document.add_paragraph("Before the figure.")
    document.add_picture(str(image(tmp_path)))
    document.add_paragraph("After the figure.")
    document.save(source)
    parts = list(visual_parts(source))
    assert parts[0]["text"] == "Before the figure."
    assert "image" in parts[1] and parts[1]["location"].startswith("block:2:image:")
    assert parts[2]["text"] == "After the figure."


def test_docx_preserves_table_cell_relationships_and_image_order(tmp_path):
    import docx
    source = tmp_path / "limits.docx"
    document = docx.Document()
    document.add_paragraph("Equipment limits.")
    table = document.add_table(rows=3, cols=2)
    for row, values in zip(table.rows, [("Component", "Limit"), ("Pump", "30 V"), ("Fan", "12 V")]):
        for cell, value in zip(row.cells, values):
            cell.text = value
    document.add_picture(str(image(tmp_path)))
    document.add_paragraph("End of limits.")
    document.save(source)
    parts = list(visual_parts(source))
    assert [part["location"] for part in parts] == ["block:1", "block:2", "block:3:image:1", "block:4"]
    assert parts[1]["text"] == "Component\tLimit\nPump\t30 V\nFan\t12 V"
    assert "image" in parts[2]
    assert parts[3]["text"] == "End of limits."


def test_docx_preserves_paragraph_tabs_breaks_and_adjacent_runs(tmp_path):
    import docx
    source = tmp_path / "instructions.docx"
    document = docx.Document()
    paragraph = document.add_paragraph()
    paragraph.add_run("Input")
    paragraph.add_run(":").add_tab()
    run = paragraph.add_run("24 V")
    run.add_break()
    paragraph.add_run("Keep dry.")
    document.save(source)
    assert list(visual_parts(source))[0]["text"] == "Input:\t24 V\nKeep dry."


def test_docx_preserves_nested_table_text_in_its_cell(tmp_path):
    import docx
    source = tmp_path / "nested.docx"
    document = docx.Document()
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Power"
    table.cell(0, 1).text = "Range"
    nested = table.cell(0, 1).add_table(rows=1, cols=2)
    nested.cell(0, 0).text = "Minimum"
    nested.cell(0, 1).text = "12 V"
    document.save(source)
    assert list(visual_parts(source))[0]["text"] == "Power\tRange\nMinimum\t12 V\n"


def test_pdf_and_image_limits_reject_oversized_sources_before_model_calls(tmp_path, monkeypatch):
    from lib.infrastructure import document_vision
    source = tmp_path / "scan.pdf"
    first = Image.new("RGB", (120, 90), "blue")
    first.save(source, save_all=True, append_images=[first])
    monkeypatch.setattr(document_vision, "MAX_PAGES", 1)
    with pytest.raises(ValueError, match="document_vision_page_limit"):
        list(visual_parts(source))
    monkeypatch.setattr(document_vision, "MAX_IMAGE_PIXELS", 100)
    with pytest.raises(ValueError, match="document_image_too_large"):
        list(visual_parts(image(tmp_path)))


def test_text_documents_in_vision_mode_remain_native_and_do_not_call_model(tmp_path):
    config(tmp_path)
    source = tmp_path / "guide.md"
    source.write_text(TEXT, encoding="utf-8")
    run_id = engine.create_run(tmp_path / "output", sources=[source], targets=["cpt"], settings_root=tmp_path,
                               document_parser=parser())
    client = Vision()
    state = engine.Workflow(tmp_path / "output", run_id, tmp_path, generator=client).execute()
    assert state["status"] == "completed"
    assert client.calls == []


@pytest.mark.parametrize("api_format", ["chat", "responses", "anthropic"])
def test_document_vision_uses_existing_protocol_adapter_and_durable_stream(tmp_path, api_format):
    helpers_path = Path(__file__).with_name("model_streaming_fixtures.py")
    spec = importlib.util.spec_from_file_location("vision_stream_helpers", helpers_path)
    helpers = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(helpers)
    config(tmp_path, api_format=api_format)
    calls = []

    def handle(request):
        calls.append(json.loads(request.content))
        return helpers.sse_response(helpers.fixture(api_format))

    client = helpers.sdk_client(api_format, handle)
    try:
        run_id = engine.create_run(tmp_path / "output", sources=[image(tmp_path)], targets=["cpt"],
                                   settings_root=tmp_path, document_parser=parser())
        run = engine.Workflow(tmp_path / "output", run_id, tmp_path, generator=client)
        part = next(visual_parts(image(tmp_path)))
        assert run.ask(["image-1", "vision"], "vision", "workflow.document_vision", {}, image=part["image"]) == {"answer": "ok"}
        request = calls[0]
        if api_format == "responses":
            assert request["input"][1]["content"][1]["type"] == "input_image"
        elif api_format == "anthropic":
            assert request["messages"][0]["content"][1]["type"] == "image"
        else:
            assert request["messages"][1]["content"][1]["type"] == "image_url"
        streams = read_streams(run.path, active=False, run_attempt=0)
        assert streams[0]["stage"] == "ingest" and streams[0]["role"] == "vision"
        assert streams[0]["text"] == '{"answer":"ok"}' and streams[0]["status"] == "completed"
    finally:
        client.client.close()
