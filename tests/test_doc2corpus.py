"""doc2corpus 离线测试：导入/清洗/分块/全局去重（fixture 文档）。"""
from __future__ import annotations

import pathlib

import pytest

from lib.doc2corpus import chunk_hash, chunk_text, clean_text, doc_to_corpus, import_text

ROOT = pathlib.Path(__file__).resolve().parent.parent
FIXTURE = ROOT / "tests" / "fixtures" / "doc_sample.md"


def test_import_and_clean_markdown():
    raw = import_text(FIXTURE)
    assert "# 第一章" in raw
    cleaned = clean_text("段落一\n\n\n\n段落二\n\n1234\n\n段落三")
    assert "段落一\n\n段落二" in cleaned
    assert "1234" in cleaned  # 数字可能是正文、编号或表格值，不能自动删去。
    assert "段落三" in cleaned

    code = clean_text("说明\n\n    value = 1234\n    return value\n")
    assert "    value = 1234\n    return value" in code
    assert "    value = 1234\n    return value" in chunk_text(code)[0]


def test_chunk_text_respects_paragraphs_and_headings():
    text = "\n\n".join(["# 标题A"] + [f"段落{i}：" + "知识" * 300 for i in range(10)])
    chunks = chunk_text(text, target_chars=2000)
    assert len(chunks) >= 2
    assert chunks[0].startswith("# 标题A")  # 标题硬边界开启首块
    joined = "".join(chunks)
    assert "段落9" in joined  # 内容完整保留


def test_chunk_text_splits_oversized_single_paragraph_without_losing_characters():
    paragraph = "".join(f"编号{i:04d}对应参数{i % 31}。" for i in range(300))
    chunks = chunk_text(paragraph, target_chars=250)
    assert len(chunks) > 1
    assert all(len(chunk) <= 250 for chunk in chunks)
    assert "".join(chunks) == paragraph
    overlapped = chunk_text(paragraph, target_chars=250, overlap=30)
    assert all(len(chunk) <= 250 for chunk in overlapped)
    assert overlapped[0][-30:] == overlapped[1][:30]
    assert overlapped[-1].endswith(paragraph[-30:])


def test_doc_to_corpus_with_global_dedup(tmp_path):
    out = tmp_path / "corpus.jsonl"
    manifest: set[str] = set()
    r1 = doc_to_corpus(FIXTURE, target_chars=800, manifest=manifest)
    assert r1["stats"]["kept"] >= 2
    # 二次处理同文档：全部命中 manifest
    r2 = doc_to_corpus(FIXTURE, target_chars=800, manifest=manifest)
    assert r2["stats"]["kept"] == 0
    assert r2["stats"]["dups"] == r1["stats"]["kept"]

    from lib.doc2corpus import write_corpus_jsonl

    n = write_corpus_jsonl(r1["entries"], out)
    assert n == len(r1["entries"])
    lines = out.read_text(encoding="utf-8").splitlines()
    assert all('"text"' in line and '"source"' in line for line in lines)


def test_chunk_hash_normalizes_whitespace():
    assert chunk_hash("你好 世界") == chunk_hash("你好\n世界")


def test_import_text_decodes_non_utf8_chinese(tmp_path):
    """GB18030 中文文档不再被 errors=replace 静默转成 U+FFFD 乱码。"""
    target = tmp_path / "gbk.txt"
    target.write_bytes("中文知识：分块与去重。".encode("gb18030"))
    raw = import_text(target)
    assert "中文知识" in raw
    assert "\ufffd" not in raw


def test_undecodable_cpt_source_is_quarantined_instead_of_replaced(tmp_path):
    from lib.infrastructure.training_workflow import Workflow, create_run

    source = tmp_path / "broken.txt"
    source.write_bytes(b"\x81")  # Incomplete in both UTF-8 and GB18030.
    with pytest.raises(UnicodeError, match="invalid_text_encoding"):
        doc_to_corpus(source)

    run_id = create_run(tmp_path, sources=[source], targets=["cpt"])
    workflow = Workflow(tmp_path, run_id, tmp_path)
    records = workflow.parse_source(workflow.recipe["sources"][0])
    assert len(records) == 1
    assert records[0]["status"] == "quarantined"
    assert records[0]["reason"] == "invalid_encoding"


def test_cpt_corpus_keeps_standalone_numbers_and_code_indentation(tmp_path):
    source = tmp_path / "manual.md"
    source.write_text("设备编号\n\n1234\n\n    value = 1234\n    return value\n", encoding="utf-8")
    output = doc_to_corpus(source)
    text = "\n\n".join(entry["text"] for entry in output["entries"])
    assert "\n\n1234\n\n" in text
    assert "    value = 1234\n    return value" in text


def test_docx_original_body_and_nested_table_order_is_kept_for_text_and_model_reading(tmp_path):
    import docx
    from lib.infrastructure.document_text import text_parts
    from lib.infrastructure.document_vision import visual_parts

    document = docx.Document()
    document.add_paragraph("Definition before table")
    cell = document.add_table(rows=1, cols=1).cell(0, 0)
    cell.text = "Cell introduction"
    cell.add_table(rows=1, cols=1).cell(0, 0).text = "Nested measurement: 12.7"
    cell.add_paragraph("Cell conclusion")
    document.add_paragraph("Interpretation after table")
    source = tmp_path / "measurements.docx"
    document.save(source)

    text = import_text(source)
    labels = ["Definition before table", "Cell introduction", "Nested measurement: 12.7",
              "Cell conclusion", "Interpretation after table"]
    assert [text.index(label) for label in labels] == sorted(text.index(label) for label in labels)
    assert next(text_parts(source))["text"] == text
    assert "\n".join(part["text"] for part in visual_parts(source)) == text

    # Old immutable runs retain their original text ordering and identities.
    legacy = import_text(source, source_processing_version=1)
    assert legacy.index("Interpretation after table") < legacy.index("Cell introduction")
    assert "Nested measurement" not in legacy


def test_long_prose_prefers_whole_sentences_and_preserves_every_character():
    source = "".join(f"Measurement {index}: the concentration is {index}.7 mg. " for index in range(30))
    chunks = chunk_text(source, target_chars=120)
    assert "".join(chunks) == source
    assert all(len(chunk) <= 120 for chunk in chunks)
    assert all(chunk.rstrip().endswith(".") for chunk in chunks[:-1])
    assert all(".7 mg." in chunk for chunk in chunks)


def test_markdown_headings_without_blank_lines_are_real_boundaries():
    source = "Introduction.\n# First chapter\nFirst evidence.\n## Second chapter\nSecond evidence.\n"
    chunks = chunk_text(source, target_chars=200)
    assert chunks == ["Introduction.\n", "# First chapter\nFirst evidence.\n",
                      "## Second chapter\nSecond evidence.\n"]
    assert "".join(chunks) == source


def test_markdown_code_fences_and_table_relationships_remain_intact():
    code = "```python\n# This is code, not a document heading\nvalue = 12.7\n\nprint(value)\n```\n"
    table = "| Measure | Value |\n| --- | ---: |\n| Concentration | 12.7 |\n| Limit | 21.4 |\n"
    source = "# Measurements\nBefore the example.\n\n" + code + "\n" + table + "\nAfter the table.\n"
    chunks = chunk_text(source, target_chars=45)
    assert "".join(chunks) == source
    assert all(chunk.strip() for chunk in chunks)
    assert any(code in chunk for chunk in chunks)
    assert any(table in chunk for chunk in chunks)
    assert sum("Concentration" in chunk for chunk in chunks) == 1


def test_chunk_separators_never_become_empty_evidence_or_lose_source_characters():
    source = ("\n\n" + "Sentence with a complete result. " * 12 + "\n\n"
              + "```text\n" + "A code example that stays whole.\n" * 8 + "```\n\n\n"
              + "| Measure | Value |\n| --- | --- |\n" + "| Sample | 12.7 |\n" * 8
              + "\n\nConcluding evidence.\n\n")
    chunks = chunk_text(source, target_chars=45)
    assert "".join(chunks) == source
    assert all(chunk.strip() for chunk in chunks)


def test_legacy_split_policy_stays_exact_for_saved_workflow_inputs():
    source = "Paragraph A.\n# Header without blank lines\nParagraph B.\n\nTail."
    assert chunk_text(source, target_chars=200, source_processing_version=1) == [source]
    assert chunk_text(source, target_chars=200) != [source]


def test_new_workflow_pins_source_processing_and_old_recipe_replays_legacy(tmp_path):
    from lib.infrastructure.training_workflow import Workflow, create_run

    source = tmp_path / "manual.md"
    source.write_text("A.\n# Heading\nB.", encoding="utf-8")
    run_id = create_run(tmp_path / "output", sources=[source], targets=["cpt"], chunk_chars=200)
    workflow = Workflow(tmp_path / "output", run_id, tmp_path)
    assert workflow.recipe["source_processing_version"] == 2
    records = workflow.parse_source(workflow.recipe["sources"][0])
    assert [record["text"] for record in records] == ["A.\n", "# Heading\nB."]

    # Simulate a historical recipe lacking the new explicit parser policy.
    workflow.recipe.pop("source_processing_version")
    legacy = workflow.parse_source(workflow.recipe["sources"][0])
    assert [record["text"] for record in legacy] == ["A.\n# Heading\nB."]
    assert legacy[0]["source_location"]["chunk"] == 0


def test_docx_expansion_is_bounded_before_document_library_reads_parts(tmp_path, monkeypatch):
    import docx
    from lib.infrastructure import docx_document

    source = tmp_path / "large-expansion.docx"
    docx.Document().save(source)
    monkeypatch.setattr(docx_document, "MAX_EXPANDED_DOCX_BYTES", 100)
    monkeypatch.setattr(docx, "Document", lambda *_: pytest.fail("unbounded package must not be opened"))
    with pytest.raises(ValueError, match="^document_docx_expansion_limit$"):
        import_text(source)


@pytest.mark.parametrize("suffix, contents, reason", [
    (".docx", b"not a document archive", "document_docx_parse_failed"),
    (".tex", br"\section{broken", "latex_unbalanced_group"),
])
def test_single_corrupt_source_is_isolated_while_valid_document_exports(tmp_path, suffix, contents, reason):
    from lib.infrastructure.training_workflow import Workflow, create_run, read_json, verify_artifacts

    broken = tmp_path / ("broken" + suffix)
    broken.write_bytes(contents)
    valid = tmp_path / "valid.txt"
    valid.write_text("Before starting the equipment, check the power cable. Keep the machine dry. "
                     "Disconnect the power before maintenance. Record the result after every inspection.",
                     encoding="utf-8")
    output = tmp_path / "output"
    run_id = create_run(output, sources=[broken, valid], targets=["cpt"])
    run = Workflow(output, run_id, tmp_path)
    state = run.execute()
    assert state["status"] == "needs_attention"
    assert verify_artifacts(run.path)["counts"]["cpt"] == 1
    rows = read_json(run.path / "input_records.json")
    invalid = next(row for row in rows if row["status"] == "quarantined")
    assert invalid["reason"] == reason
    assert invalid["source_location"] == {"file": broken.name, "record": "document"}
    assert any(row.get("text") == valid.read_text(encoding="utf-8") for row in rows)


def test_invalid_docx_xml_is_reported_as_a_document_error(tmp_path):
    import docx
    from zipfile import ZipFile

    source = tmp_path / "valid.docx"
    docx.Document().save(source)
    broken = tmp_path / "bad-xml.docx"
    with ZipFile(source) as original, ZipFile(broken, "w") as changed:
        for name in original.namelist():
            changed.writestr(name, b"<w:document>" if name == "word/document.xml" else original.read(name))
    with pytest.raises(ValueError, match="^document_docx_parse_failed$"):
        import_text(broken)


def test_unexpected_parser_failure_is_not_disguised_as_a_bad_document(tmp_path, monkeypatch):
    from lib.infrastructure import training_workflow as engine

    source = tmp_path / "valid.txt"
    source.write_text("Reliable source text.", encoding="utf-8")
    output = tmp_path / "output"
    run_id = engine.create_run(output, sources=[source], targets=["cpt"])
    run = engine.Workflow(output, run_id, tmp_path)

    def broken_parser(*args, **kwargs):
        raise RuntimeError("unexpected implementation failure")

    monkeypatch.setattr(engine, "import_text", broken_parser)
    with pytest.raises(RuntimeError, match="unexpected implementation failure"):
        run.parse_source(run.recipe["sources"][0])
